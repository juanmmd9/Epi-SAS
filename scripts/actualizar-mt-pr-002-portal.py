# -*- coding: utf-8 -*-
"""Genera MT-PR-002 V6 alineado al Portal de Mantenimiento EPI (GC-RE-027)."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

DESTINO = Path(
    r"C:\Users\Ing. Mecanico\Desktop\MANTENIMIENTO\PROCEDIMIENTOS E INSTRUCTIVOS"
    r"\PROCEDIMIENTOS\MT-PR-002 Mantenimiento correctivo V6.docx"
)


def set_run_font(run, size=11, bold=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold


def add_heading_epi(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 12, True)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)


def add_body(doc: Document, text: str, bold: bool = False) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11, bold)
    p.paragraph_format.space_after = Pt(4)


def add_bullet(doc: Document, text: str) -> None:
    p = doc.add_paragraph(style="List Bullet")
    # limpiar y escribir con fuente
    p.clear()
    run = p.add_run(text)
    set_run_font(run, 11)


def fill_header_row(row, headers: list[str]) -> None:
    for i, h in enumerate(headers):
        row.cells[i].text = ""
        p = row.cells[i].paragraphs[0]
        run = p.add_run(h)
        set_run_font(run, 10, True)


def add_step_row(table, no: str, actividad: str, responsable: str, obs: str, registro: str) -> None:
    row = table.add_row()
    values = [no, actividad, responsable, obs, registro]
    for i, val in enumerate(values):
        row.cells[i].text = ""
        p = row.cells[i].paragraphs[0]
        run = p.add_run(val)
        set_run_font(run, 9)


def main() -> None:
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(1.8)
        section.right_margin = Cm(1.8)

    titulo = doc.add_paragraph()
    titulo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = titulo.add_run("MT-PR-002 — PROCEDIMIENTO DE MANTENIMIENTO CORRECTIVO")
    set_run_font(r, 14, True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta.add_run("Versión 6  |  Septiembre 2026  |  Proceso: Mantenimiento  |  E.P.I. S.A.S.")
    set_run_font(r, 10)

    add_heading_epi(doc, "1. OBJETIVO")
    add_body(
        doc,
        "Establecer los pasos para asistir y reparar de manera oportuna el equipamiento "
        "e infraestructura de E.P.I. S.A.S. cuando se detecte o presente una falla no "
        "prevista, utilizando el Portal de Mantenimiento EPI como herramienta oficial "
        "de solicitud, atención, registro y seguimiento, en cumplimiento de NTC-ISO "
        "9001:2015 (procesos) y NTC-ISO/IEC 17025:2017 (Laboratorio).",
    )

    add_heading_epi(doc, "2. ALCANCE")
    add_body(
        doc,
        "Aplica a maquinaria y equipos de E.P.I. S.A.S. que presenten fallas no previstas "
        "(mecánico, eléctrico, neumático, hidráulico, electrónico e informático), en "
        "procesos productivos y Laboratorio. Involucra al personal de mantenimiento, "
        "operarios, líderes de área (solicitantes del Portal) y roles de consulta/admin.",
    )

    add_heading_epi(doc, "3. RESPONSABILIDADES")
    add_body(
        doc,
        "Coordinador de Mantenimiento: Supervisar la atención de solicitudes en el Portal, "
        "asignar técnicos cuando aplique, controlar tiempos de respuesta, compras de "
        "repuestos y el cumplimiento de registros/indicadores.",
    )
    add_body(
        doc,
        "Auxiliares / operadores de mantenimiento: Tomar o recibir solicitudes en el Portal, "
        "diagnosticar, reparar, registrar avances (incl. espera de repuesto), cerrar la "
        "solicitud y dejar evidencia del trabajo realizado.",
    )
    add_body(
        doc,
        "Líder de área / solicitante (usuario Portal): Reportar la falla en el módulo de "
        "Solicitudes del área, colocar aviso de fuera de servicio cuando corresponda y "
        "confirmar la recepción del equipo a satisfacción.",
    )
    add_body(
        doc,
        "Operario del equipo: Detectar la falla, asegurar el equipo (fuera de servicio) y "
        "apoyar la creación de la solicitud en el Portal junto con el líder de área.",
    )

    add_heading_epi(doc, "4. DEFINICIONES")
    add_body(
        doc,
        "Mantenimiento correctivo: Conjunto de actividades para corregir fallas y "
        "anormalidades en equipos a medida que se presentan, con la máquina fuera de "
        "servicio, hasta restablecer su operación segura y conforme.",
    )
    add_body(
        doc,
        "Portal de Mantenimiento EPI: Aplicación web oficial (GC-RE-027) para programar, "
        "solicitar, registrar y consultar mantenimientos preventivos/correctivos y "
        "solicitudes de área; sustituye Excel/papel como fuente oficial tras la fecha de corte.",
    )
    add_body(
        doc,
        "Solicitud de área: Registro digital en el Portal mediante el cual el área reporta "
        "una falla o necesidad de intervención correctiva.",
    )

    add_heading_epi(doc, "5. REQUISITOS")
    add_body(doc, "Para los mantenimientos correctivos se debe contar con:")
    for item in [
        "Catálogos de equipos, manuales del fabricante, planos.",
        "Experiencia y competencia de técnicos (matriz de conocimientos del Portal).",
        "Acceso al Portal (usuario y rol: solicitante, operador, admin o consulta).",
        "Repuestos y herramientas según la falla; compras conforme a CO-PR-001 / Drive Compras.",
    ]:
        add_bullet(doc, item)

    add_heading_epi(doc, "6. CONTENIDO — CONDICIONES GENERALES")
    add_body(
        doc,
        "Identificar el equipo fuera de servicio que presente falla no prevista. Colocar "
        "aviso de fuera de servicio (MT-RE-011 o equivalente) y reportar de inmediato en "
        "el Portal. Se mantienen como referencia las partes típicas a revisar "
        "(neumático, hidráulico, mecánico, eléctrico) y el listado de herramientas del "
        "procedimiento anterior.",
    )

    add_heading_epi(doc, "7. MANTENIMIENTO CORRECTIVO — ACTIVIDADES")
    table = doc.add_table(rows=1, cols=5)
    table.style = "Table Grid"
    fill_header_row(
        table.rows[0],
        ["No.", "ACTIVIDADES", "RESPONSABLE", "OBSERVACIONES", "REGISTROS"],
    )

    pasos = [
        (
            "1",
            "Detectar la falla en el equipo al no poder ejecutar correctamente una operación.",
            "Persona que opera el equipo",
            "Colocar aviso «Fuera de servicio». Asegurar el área y detener uso del equipo.",
            "MT-RE-011\nAviso fuera de servicio",
        ),
        (
            "2",
            "Crear la solicitud de mantenimiento en el Portal (módulo Solicitudes del área), "
            "describiendo la falla y el equipo/locación.",
            "Líder de área o persona que opera el equipo (usuario solicitante)",
            "En lo posible indicar si la falla es mecánica, neumática, eléctrica o hidráulica. "
            "El Portal es la vía oficial (reemplaza la solicitud solo por escrito/papel).",
            "Portal — Solicitudes de área",
        ),
        (
            "3",
            "Recibir la solicitud en el Portal, dar respuesta y realizar diagnóstico.",
            "Coordinador / auxiliar de mantenimiento",
            "Meta de respuesta: 10 minutos (según acuerdo operativo). "
            "Tomar la solicitud de la bandeja o recibir asignación del coordinador.",
            "Portal — Bandeja / Mis solicitudes\nAsignación correctivo",
        ),
        (
            "4",
            "Realizar la reparación de la máquina o equipo.",
            "Coordinador / auxiliar de mantenimiento (o servicio contratado)",
            "Puede ser personal interno o servicio externo. Registrar avances en el Portal.",
            "Portal — Correctivo / Solicitud",
        ),
        (
            "5",
            "Definir necesidad de repuestos y listar piezas o mecanismos a cambiar.",
            "Auxiliar de mantenimiento",
            "Si falta repuesto, marcar en el Portal el estado de espera de repuesto.",
            "Portal — Solicitud (espera repuesto)\nListado de repuestos",
        ),
        (
            "6",
            "Solicitar los repuestos al departamento de Compras (CO-PR-001).",
            "Coordinador de mantenimiento",
            "Ser explícito en referencias y medidas.",
            "Drive Compras\nCO-PR-001",
        ),
        (
            "7",
            "Instalar los repuestos recibidos y continuar la reparación.",
            "Auxiliar de mantenimiento",
            "Verificar lo recibido frente a lo solicitado. Quitar espera de repuesto en el Portal al reanudar.",
            "Drive Compras\nPortal — Solicitud",
        ),
        (
            "8",
            "Ensayar / probar el equipo en reparación.",
            "Auxiliar de mantenimiento y operario del equipo",
            "Repetir las pruebas cuantas veces sea necesario para asegurar el correcto funcionamiento.",
            "Portal — Solicitud / Correctivo",
        ),
        (
            "9",
            "Evaluar necesidad de calibración; si aplica, solicitar el servicio antes de liberar.",
            "Coordinador de mantenimiento",
            "Si requiere calibración, no regresar a servicio hasta ejecutarla.",
            "Solicitud de calibración\nPortal (nota en solicitud)",
        ),
        (
            "10",
            "Retirar el aviso de fuera de servicio y cerrar la solicitud en el Portal.",
            "Personal de mantenimiento",
            "Solo si el equipo fue probado y funciona correctamente. Completar datos de cierre "
            "(actividad, técnicos, tiempos).",
            "Portal — Cierre de solicitud\nMT-RE-011 (retiro aviso)",
        ),
        (
            "11",
            "Registrar el mantenimiento correctivo realizado (trazabilidad e indicadores).",
            "Auxiliar / Coordinador de mantenimiento",
            "El registro oficial queda en el Portal (ya no se entregan a fin de mes formatos "
            "MT-RE-001 a 004 en papel como fuente primaria). Los indicadores se toman del Portal.",
            "Portal — Correctivo / Indicadores\nMT-RE-045 (si se usó en la intervención)",
        ),
        (
            "12",
            "Entregar el equipo a satisfacción con confirmación del líder o colaborador del área.",
            "Líder de área / operario y mantenimiento",
            "El área confirma que la acción quedó a satisfacción. Laboratorio: entregar informe "
            "con las reparaciones realizadas.",
            "Portal — Solicitud cerrada\nInforme (Laboratorio)",
        ),
        (
            "13",
            "Solo Laboratorio: actualizar hoja de vida y matriz de equipamiento.",
            "Director / personal de Laboratorio",
            "Actualizar información del equipamiento tras la intervención.",
            "LA-RE-055\nLA-RE-015\nPortal — Hojas (si aplica)",
        ),
    ]
    for paso in pasos:
        add_step_row(table, *paso)

    add_heading_epi(doc, "8. DOCUMENTOS A CONSULTAR")
    for item in [
        "Portal de Mantenimiento EPI (fuente oficial de solicitudes y registros correctivos)",
        "GC-RE-027 Gestión del cambio (implementación del Portal)",
        "CO-PR-001 Procedimiento de compras",
        "Catálogos, manuales y planos del fabricante",
        "MT-PR-001 Procedimiento de mantenimiento preventivo",
    ]:
        add_bullet(doc, item)

    add_heading_epi(doc, "9. DOCUMENTOS A DILIGENCIAR / REGISTROS")
    for item in [
        "Portal — Solicitudes de área / Correctivo (registro oficial)",
        "MT-RE-011 Aviso fuera de servicio",
        "Drive Compras (solicitud de repuestos)",
        "MT-RE-045 (cuando la intervención correctiva se documente en ese formato)",
        "LA-RE-055 / LA-RE-015 (Laboratorio)",
    ]:
        add_bullet(doc, item)

    add_heading_epi(doc, "10. SALIDAS NO CONFORMES (resumen)")
    add_body(
        doc,
        "Se mantienen los controles del proceso alineados a ISO 9001: no liberar equipos "
        "sin condiciones técnicas, sin verificación funcional, sin registro completo en el "
        "Portal, ni con personal no competente. Política: sin registro en Portal = sin "
        "liberación. Reportar al Coordinador de Mantenimiento / SGC.",
    )

    add_heading_epi(doc, "11. REVISIÓN Y APROBACIÓN")
    firmas = doc.add_table(rows=2, cols=3)
    firmas.style = "Table Grid"
    fill_header_row(firmas.rows[0], ["ELABORÓ", "REVISÓ", "APROBÓ"])
    firmas.rows[1].cells[0].text = "JUAN MANUEL MONCAYO\nCoordinador de Mantenimiento"
    firmas.rows[1].cells[1].text = "MIGUEL TORRES\nDirector de Operaciones"
    firmas.rows[1].cells[2].text = "SANDRA ANGEL\nDirectora SGC"
    for cell in firmas.rows[1].cells:
        for p in cell.paragraphs:
            for run in p.runs:
                set_run_font(run, 9)

    add_heading_epi(doc, "12. CONTROL DE CAMBIOS")
    hist = doc.add_table(rows=1, cols=3)
    hist.style = "Table Grid"
    fill_header_row(hist.rows[0], ["VERSIÓN", "FECHA", "MOTIVO"])
    historial = [
        ("2", "Septiembre 2021", "Ajuste de cargos y registros de solicitud/compras."),
        ("3", "Septiembre 2022", "Ajustes de forma por laboratorio; Drive y LA-RE-015/055."),
        ("4", "Febrero 2023 / Enero 2026", "Definiciones de tipos de mantenimiento; firmas."),
        (
            "6",
            "Septiembre 2026",
            "Alineación al Portal de Mantenimiento EPI (GC-RE-027): la solicitud, atención, "
            "espera de repuesto, cierre e indicadores del correctivo se gestionan en el Portal "
            "como fuente oficial; se actualizan responsabilidades (solicitante/operador), "
            "registros y documentos a consultar/diligenciar. Se deja de usar MT-RE-001 a 004 "
            "en papel como registro primario.",
        ),
    ]
    for ver, fecha, motivo in historial:
        row = hist.add_row()
        for i, val in enumerate([ver, fecha, motivo]):
            row.cells[i].text = ""
            p = row.cells[i].paragraphs[0]
            run = p.add_run(val)
            set_run_font(run, 9)

    DESTINO.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(DESTINO))
    print(f"OK -> {DESTINO}")


if __name__ == "__main__":
    main()
